"""
generate_vols_livrable.py
Prépare le livrable de programmes de vols pour import Philippe (table ENPA_Vol_Tmp).

Contrairement aux trains (données GTFS brutes à traiter), le fichier de vols fourni
par Nicolas est DÉJÀ au format d'import : 32 colonnes identiques à ENPA_Vol_Tmp,
mêmes en-têtes que le fichier mai-juin précédemment importé.

Ce script se contente donc de :
  1. Conserver à l'identique les 32 en-têtes et les valeurs/types de cellules
     (l'import dans ENPA_Vol_Tmp se fait par position — ne PAS renommer les colonnes).
  2. Filtrer les lignes dont la fenêtre de validité (Effective From/To) intersecte
     réellement la période juillet → septembre 2026.
  3. Ajouter un onglet 'Controle' (synthèse par aéroport / faisceau EVER + dates).

Usage :
  python scripts/generate_vols_livrable.py
"""

import datetime
import os
import sys

import openpyxl

# ── Paramètres ────────────────────────────────────────────────────────────────
SRC    = r'Z:\Transfert Info\Nicolas\EVER_2026\4_Programme_vols\EVER_Programmes_vols_juil_aout_sept_20260609.xlsx'
OUTPUT = r'C:\Users\y_bicrel\Downloads\EVER_Programmes_vols_juil_aout_sept_20260609_livrable.xlsx'
SHEET  = 'VF'

PERIODE_MIN = datetime.date(2026, 7, 1)
PERIODE_MAX = datetime.date(2026, 9, 30)

COL_EFF_FROM = 'Effective From'
COL_EFF_TO   = 'Effective To'
COL_DEP_APT  = 'Dep Airport Code'
COL_FAISCEAU = 'Faisceau EVER'


def to_date(x):
    if isinstance(x, datetime.datetime):
        return x.date()
    if isinstance(x, datetime.date):
        return x
    return None


def main():
    if not os.path.exists(SRC):
        print(f'ERREUR : fichier source introuvable : {SRC}')
        sys.exit(1)

    print(f'Lecture de {SRC}...')
    wb_src = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws_src = wb_src[SHEET]

    rows = list(ws_src.iter_rows(values_only=True))
    header = list(rows[0])
    data   = rows[1:]
    print(f'  {len(data)} lignes lues, {len(header)} colonnes')

    idx = {h: i for i, h in enumerate(header)}
    i_from = idx[COL_EFF_FROM]
    i_to   = idx[COL_EFF_TO]
    i_apt  = idx[COL_DEP_APT]
    i_fais = idx[COL_FAISCEAU]

    # ── Filtre période ──
    kept = []
    dropped = []
    no_dates = []
    for r in data:
        f = to_date(r[i_from])
        t = to_date(r[i_to])
        if f is None or t is None:
            no_dates.append(r)
            kept.append(r)   # on conserve les lignes sans dates (prudence)
            continue
        # Le vol opère dans la période si sa fenêtre de validité l'intersecte
        if f <= PERIODE_MAX and t >= PERIODE_MIN:
            kept.append(r)
        else:
            dropped.append(r)

    print(f'  Conservées (intersectent juil-sept) : {len(kept)}')
    print(f'  Retirées (hors période)             : {len(dropped)}')
    if no_dates:
        print(f'  Sans dates (conservées par prudence): {len(no_dates)}')

    # ── Écriture ──
    print(f'\nÉcriture de {OUTPUT}...')
    wb_out = openpyxl.Workbook()
    ws = wb_out.active
    ws.title = SHEET

    # En-têtes à l'identique
    ws.append(header)
    # Données (valeurs et types préservés)
    for r in kept:
        ws.append(list(r))

    # ── Onglet contrôle ──
    ws_ctrl = wb_out.create_sheet('Controle')

    # Synthèse par aéroport de départ enquêté
    from collections import Counter
    by_apt  = Counter(r[i_apt] for r in kept)
    by_fais = Counter(r[i_fais] for r in kept)

    # Plage de dates effective
    all_from = [to_date(r[i_from]) for r in kept if to_date(r[i_from])]
    all_to   = [to_date(r[i_to])   for r in kept if to_date(r[i_to])]

    ws_ctrl.append(['Indicateur', 'Valeur'])
    ws_ctrl.append(['Lignes livrées', len(kept)])
    ws_ctrl.append(['Lignes retirées (hors juil-sept)', len(dropped)])
    ws_ctrl.append(['Période ciblée', f'{PERIODE_MIN} → {PERIODE_MAX}'])
    ws_ctrl.append(['Effective From min', str(min(all_from)) if all_from else ''])
    ws_ctrl.append(['Effective To max',   str(max(all_to))   if all_to   else ''])
    ws_ctrl.append([])
    ws_ctrl.append(['Aéroport de départ', 'Nbre de vols'])
    for apt, n in sorted(by_apt.items(), key=lambda kv: -kv[1]):
        ws_ctrl.append([apt, n])
    ws_ctrl.append([])
    ws_ctrl.append(['Faisceau EVER', 'Nbre de vols'])
    for f, n in sorted(by_fais.items(), key=lambda kv: -kv[1]):
        ws_ctrl.append([f, n])

    wb_out.save(OUTPUT)
    print(f'Fichier écrit : {OUTPUT}')
    print(f'  Onglet {SHEET}       : {len(kept)} lignes ({len(header)} colonnes)')
    print(f'  Onglet Controle  : synthèse')

    # Récap console
    print('\nRépartition par aéroport de départ :')
    for apt, n in sorted(by_apt.items(), key=lambda kv: -kv[1]):
        print(f'  {apt:6s} {n:6d}')


if __name__ == '__main__':
    main()
